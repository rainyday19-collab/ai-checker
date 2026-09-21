from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfWriter

import database
from ai_grading import build_model_input
from assessment_service import mock_grade_assessment
from document_processing import MAX_FILE_BYTES, SUPPORTED_TYPES, extract_content
from grading import prepare_grading_input
from main import app
from mark_scheme_storage import managed_path


DOCX_MIME = SUPPORTED_TYPES[".docx"]


def image_bytes(format_name="PNG", color="navy"):
    output = BytesIO()
    Image.new("RGB", (100, 80), color).save(output, format=format_name)
    return output.getvalue()


def docx_bytes(*, paragraphs=(), table=None, images=()):
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                word_table.cell(row_index, cell_index).text = value
    for image in images:
        document.add_picture(BytesIO(image), width=Inches(1))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def with_archive_members(data, members):
    output = BytesIO(data)
    with ZipFile(output, "a", ZIP_DEFLATED) as archive:
        for name, content in members:
            archive.writestr(name, content)
    return output.getvalue()


def blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class DocxProcessingTests(unittest.TestCase):
    def test_headings_and_lists_keep_semantic_boundaries(self):
        document = Document()
        document.add_heading("Section one", level=1)
        document.add_paragraph("First point", style="List Bullet")
        output = BytesIO()
        document.save(output)
        processed = extract_content(output.getvalue(), "structured.docx", DOCX_MIME, "Student work")
        self.assertIn("[Heading]\nSection one", processed.text)
        self.assertIn("[List Item]\nFirst point", processed.text)

    def test_paragraphs_and_table_preserve_document_order(self):
        data = docx_bytes(paragraphs=("Question 1 answer",), table=(("Criterion", "Answer"), ("Method", "Correct")))
        # Add a paragraph after the table to exercise interleaved body order.
        document = Document(BytesIO(data))
        document.add_paragraph("Answer after table")
        output = BytesIO()
        document.save(output)
        processed = extract_content(output.getvalue(), "answers.docx", DOCX_MIME, "Student work")
        self.assertEqual(processed.type, "docx")
        self.assertLess(processed.text.index("Question 1 answer"), processed.text.index("[Table]"))
        self.assertLess(processed.text.index("Row 2: Method | Correct"), processed.text.index("Answer after table"))
        self.assertIn("[Paragraph]", processed.text)

    def test_embedded_png_and_jpeg_are_in_one_model_input(self):
        data = docx_bytes(paragraphs=("Use both diagrams as evidence.",),
                          images=(image_bytes("PNG", "navy"), image_bytes("JPEG", "gold")))
        processed = extract_content(data, "illustrated.docx", DOCX_MIME, "Student work")
        self.assertEqual([image.content_type for image in processed.embedded_images], ["image/png", "image/jpeg"])
        prepared = prepare_grading_input(processed, criteria_text="Award four marks.")
        self.assertEqual(prepared.student_work.kind, "docx")
        parts = build_model_input(prepared)[0]["content"]
        self.assertEqual(sum(part["type"] == "input_image" for part in parts), 2)
        self.assertEqual(sum("DOCX EMBEDDED IMAGE" in part.get("text", "") for part in parts), 2)

    def test_archive_and_extracted_content_limits(self):
        base = docx_bytes(paragraphs=("Answer",))
        too_many = with_archive_members(base, ((f"custom/item-{index}.txt", b"x") for index in range(1001)))
        traversal = with_archive_members(base, [("../outside.txt", b"x")])
        huge_text = docx_bytes(paragraphs=("x" * 500_001,))
        for data, expected in [(too_many, 413), (traversal, 422), (huge_text, 413)]:
            with self.subTest(expected=expected):
                with self.assertRaises(Exception) as raised:
                    extract_content(data, "unsafe.docx", DOCX_MIME, "Student work")
                self.assertEqual(raised.exception.status_code, expected)

    def test_excessive_embedded_image_count_is_rejected(self):
        images = tuple(image_bytes("PNG", (index, index * 3 % 256, index * 7 % 256)) for index in range(21))
        data = docx_bytes(paragraphs=("Illustrated answer",), images=images)
        with self.assertRaises(Exception) as raised:
            extract_content(data, "too-many-images.docx", DOCX_MIME, "Student work")
        self.assertEqual(raised.exception.status_code, 413)
        self.assertIn("at most 20", raised.exception.detail)


class DocxEndpointTests(unittest.TestCase):
    def setUp(self):
        temporary = self.enterContext(TemporaryDirectory())
        self.root = Path(temporary)
        self.enterContext(patch("database.DB_PATH", self.root / "test.db"))
        self.enterContext(patch.dict("os.environ", {"GRADING_MODE": "mock"}, clear=True))
        self.enterContext(patch("assessment_service.load_dotenv"))
        self.enterContext(patch("ai_grading.load_dotenv"))
        self.ai_client = self.enterContext(patch("ai_grading.AsyncOpenAI"))
        self.client = self.enterContext(TestClient(app))
        self.parent = self.client.post("/api/classes", json={"name": "Test", "subject": "CS"}).json()["id"]
        self.assignment = self.client.post(f"/api/classes/{self.parent}/assignments", json={
            "title": "DOCX assignment", "mark_scheme_text": "Award four marks.",
        }).json()["id"]
        self.png = image_bytes()
        self.docx = docx_bytes(paragraphs=("Student answer in Word.",),
                               table=(("Question", "Answer"), ("1", "Four")), images=(self.png,))

    def tearDown(self):
        self.ai_client.assert_not_called()

    def grade(self, uploads, assignment=None):
        return self.client.post(f"/api/assignments/{assignment or self.assignment}/submissions/grade",
            data={"student_name": "DOCX student"}, files=[("student_work", upload) for upload in uploads])

    def test_docx_student_mock_grade_is_one_result_without_binary_persistence(self):
        with patch("assessment_service.mock_grade_assessment", wraps=mock_grade_assessment) as grader:
            response = self.grade([("answers.docx", self.docx, DOCX_MIME)])
        self.assertEqual(response.status_code, 201)
        saved = response.json()
        self.assertEqual(saved["original_filenames"], ["answers.docx"])
        self.assertEqual(saved["page_count"], 1)
        grader.assert_awaited_once()
        prepared = grader.call_args.args[0]
        self.assertEqual(prepared.student_work.kind, "docx")
        self.assertIn("Student answer in Word.", prepared.student_work.text)
        self.assertEqual(len(prepared.student_work.embedded_images), 1)
        self.assertEqual(len(self.client.get("/api/assessments").json()), 1)
        self.assertEqual(len(self.client.get(f"/api/assignments/{self.assignment}/submissions").json()), 1)
        with database.connection() as connection:
            rows = [dict(row) for table in ("submissions", "assessments")
                    for row in connection.execute(f"SELECT * FROM {table}")]
        self.assertFalse(any(isinstance(value, bytes) for row in rows for value in row.values()))
        self.assertNotIn("Student answer in Word.", str(rows))

    def test_new_assessment_accepts_docx(self):
        response = self.client.post("/api/assess", data={"criteria_text": "Award four marks."},
            files={"student_work": ("answers.docx", self.docx, DOCX_MIME)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["grading_mode"], "mock")

    def test_docx_mark_scheme_is_private_and_used(self):
        scheme = docx_bytes(paragraphs=("Award four marks for the correct method.",),
                            table=(("Criterion", "Marks"), ("Method", "4")), images=(self.png,))
        response = self.client.post(f"/api/classes/{self.parent}/assignments", data={"title": "Word scheme"},
            files={"mark_scheme_file": ("scheme.docx", scheme, DOCX_MIME)})
        self.assertEqual(response.status_code, 201)
        values = response.json()
        self.assertEqual(values["mark_scheme"]["type"], "docx")
        path = managed_path(database.get_assignment(values["id"])["mark_scheme_key"])
        self.assertEqual(path.read_bytes(), scheme)
        self.assertNotIn(str(path), response.text)
        with patch("assessment_service.mock_grade_assessment", wraps=mock_grade_assessment) as grader:
            self.assertEqual(self.grade([("page.png", self.png, "image/png")], values["id"]).status_code, 201)
        prepared = grader.call_args.args[0]
        self.assertEqual(prepared.mark_scheme.kind, "docx")
        self.assertIn("Row 2: Method | 4", prepared.mark_scheme.text)
        self.assertEqual(len(prepared.mark_scheme.embedded_images), 1)

    def test_invalid_docx_combinations_and_legacy_doc(self):
        cases = [
            ([("answers.docx", self.docx, DOCX_MIME), ("page.png", self.png, "image/png")], 422),
            ([("answers.docx", self.docx, DOCX_MIME), ("work.pdf", blank_pdf(), "application/pdf")], 422),
            ([("one.docx", self.docx, DOCX_MIME), ("two.docx", self.docx, DOCX_MIME)], 422),
            ([("legacy.doc", self.docx, "application/msword")], 415),
        ]
        for uploads, expected in cases:
            with self.subTest(expected=expected, names=[upload[0] for upload in uploads]):
                self.assertEqual(self.grade(uploads).status_code, expected)
        self.assertEqual(self.client.get("/api/assessments").json(), [])

    def test_corrupt_renamed_and_oversized_docx_are_rejected(self):
        fake_zip = BytesIO()
        with ZipFile(fake_zip, "w", ZIP_DEFLATED) as archive:
            archive.writestr("notes.txt", "not a Word document")
        cases = [
            (("corrupt.docx", b"not a zip", DOCX_MIME), 422),
            (("renamed.docx", fake_zip.getvalue(), DOCX_MIME), 422),
            (("large.docx", b"x" * (MAX_FILE_BYTES + 1), DOCX_MIME), 413),
        ]
        for upload, expected in cases:
            with self.subTest(name=upload[0]):
                response = self.grade([upload])
                self.assertEqual(response.status_code, expected)
                self.assertNotIn(str(self.root), response.text)


if __name__ == "__main__":
    unittest.main()
